from __future__ import annotations

from pathlib import Path

import pytest

from navjev.build.parsers import (
    Heading,
    MarkdownParser,
    NoStructureError,
    PdfParser,
    assemble_tree,
    markdown_headings,
    parse,
)


def test_markdown_ignores_headings_inside_fences(fixtures_dir: Path) -> None:
    text = (fixtures_dir / "sample.md").read_text()
    titles = [t for _, _, t in markdown_headings(text)]
    assert "not a heading, this is a comment in code" not in titles
    assert "tilde fences too" not in titles
    assert titles == [
        "Acme Corp Annual Report 2023",
        "Business",
        "Segments",
        "Financial Statements",
        "Cash Flow",
        "Balance Sheet",
        "Executive Compensation",
        "Deep jump under compensation",
    ]


def test_markdown_tree_shape(fixtures_dir: Path) -> None:
    tree = MarkdownParser().parse(fixtures_dir / "sample.md")
    assert tree.structure_source == "native"
    assert tree.title == "Acme Corp Annual Report 2023"  # sole H1 becomes the root
    assert "Preamble paragraph" in tree.root.text
    assert [c.title for c in tree.children_of(tree.root_id)] == [
        "Business",
        "Financial Statements",
        "Executive Compensation",
    ]
    fin = next(n for n in tree.nodes.values() if n.title == "Financial Statements")
    assert [c.title for c in tree.children_of(fin.id)] == ["Cash Flow", "Balance Sheet"]
    assert fin.text == "Consolidated statements follow."  # own text only
    cash = next(n for n in tree.nodes.values() if n.title == "Cash Flow")
    assert "$412 million" in cash.text and cash.depth == 2
    deep = next(n for n in tree.nodes.values() if n.title == "Deep jump under compensation")
    assert tree.path_to(deep.id) == [
        "Acme Corp Annual Report 2023",
        "Executive Compensation",
        "Deep jump under compensation",
    ]
    assert deep.depth == 2  # nesting is by relative level, not by heading count
    assert cash.char_span is not None
    text = (fixtures_dir / "sample.md").read_text()
    assert text[cash.char_span[0] : cash.char_span[1]].startswith("### Cash Flow")


def test_markdown_without_headings_is_refused(tmp_path: Path) -> None:
    p = tmp_path / "flat.md"
    p.write_text("just prose\nno headings\n")
    with pytest.raises(NoStructureError):
        MarkdownParser().parse(p)


def test_multiple_h1_get_a_synthetic_root() -> None:
    tree = MarkdownParser().parse_text("# A\n\ntext a\n\n# B\n\ntext b\n", "d", "Doc")
    assert tree.title == "Doc"
    assert [c.title for c in tree.children_of("root")] == ["A", "B"]


def test_assemble_tree_nests_by_nearest_smaller_level() -> None:
    headings = [
        Heading(1, "One", ""),
        Heading(3, "One.deep", ""),
        Heading(2, "One.two", ""),
        Heading(1, "Two", ""),
    ]
    tree = assemble_tree("d", "Doc", "", headings, "", "test", "native")
    one = tree.nodes["n0001"]
    assert [tree.nodes[c].title for c in one.children] == ["One.deep", "One.two"]
    assert tree.nodes["n0002"].depth == 2 and tree.nodes["n0003"].depth == 2
    assert [tree.nodes[c].title for c in tree.root.children] == ["One", "Two"]


def _make_pdf(path: Path, with_outline: bool) -> None:
    import fitz

    doc = fitz.open()
    pages = [
        "Cover page of the report",
        "Business\nAcme makes widgets.",
        "Financial Statements\nConsolidated statements follow.",
        "Cash Flow\nCapital expenditure was $412 million.",
        "Balance Sheet\nTotal assets of $9.1 billion.",
    ]
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    if with_outline:
        doc.set_toc(
            [
                [1, "Business", 2],
                [1, "Financial Statements", 3],
                [2, "Cash Flow", 4],
                [2, "Balance Sheet", 5],
            ]
        )
    doc.save(str(path))
    doc.close()


def test_pdf_outline_gives_page_spans(tmp_path: Path) -> None:
    p = tmp_path / "report.pdf"
    _make_pdf(p, with_outline=True)
    tree = PdfParser().parse(p)
    assert tree.structure_source == "native"
    assert tree.root.page_span == (1, 1)
    titles = {n.title: n for n in tree.nodes.values()}
    assert titles["Financial Statements"].page_span == (3, 3)
    assert titles["Cash Flow"].page_span == (4, 4)
    assert titles["Balance Sheet"].page_span == (5, 5)
    assert "$412 million" in titles["Cash Flow"].text
    assert [tree.nodes[c].title for c in titles["Financial Statements"].children] == [
        "Cash Flow",
        "Balance Sheet",
    ]


def test_pdf_without_outline_is_refused_unless_fallback_given(tmp_path: Path) -> None:
    p = tmp_path / "flat.pdf"
    _make_pdf(p, with_outline=False)
    with pytest.raises(NoStructureError, match="no outline"):
        parse(p)


def test_dispatch_rejects_unknown_extension(tmp_path: Path) -> None:
    p = tmp_path / "x.xyz"
    p.write_text("")
    with pytest.raises(ValueError):
        parse(p)


def test_docx_heading_styles(tmp_path: Path) -> None:
    import docx

    from navjev.build.parsers import DocxParser

    d = docx.Document()
    d.add_heading("Report", level=1)
    d.add_paragraph("intro")
    d.add_heading("Business", level=2)
    d.add_paragraph("widgets")
    d.add_heading("Financials", level=2)
    d.add_paragraph("numbers")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "capex"
    table.rows[0].cells[1].text = "412"
    p = tmp_path / "r.docx"
    d.save(str(p))
    tree = DocxParser().parse(p)
    assert tree.title == "Report"
    assert [c.title for c in tree.children_of("root")] == ["Business", "Financials"]
    fin = next(n for n in tree.nodes.values() if n.title == "Financials")
    assert "numbers" in fin.text and "capex\t412" in fin.text
